from __future__ import annotations

from collections.abc import Iterable
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.enums import CallType, ImportStatus, PTOType, PriorityLevel, TaskStatus
from app.core.tracker_rules import is_wireless_reference_row
from app.imports.base import (
    ImportContext,
    find_header_row,
    infer_status,
    is_site_like,
    normalize_header,
    normalize_site_name,
    parse_date_cell,
    parse_pto_range,
    parse_tminus_due,
)
from app.models.imports import Import
from app.models.people import Call, PTO
from app.models.planning import Site


def _workstream_key_for_sheet(sheet_name: str) -> str:
    if "wireless" in sheet_name.lower():
        return "WIRELESS"
    if "sd-wan" in sheet_name.lower():
        return "SDWAN"
    if "sda" in sheet_name.lower():
        return "SDA"
    return "PROGRAM"


def parse_ops_team_sheet(ctx: ImportContext, ws) -> None:
    for row_idx in range(3, ws.max_row + 1):
        name = ws.cell(row_idx, 3).value
        email = ws.cell(row_idx, 5).value
        if not name or not email:
            continue
        user = ctx.resolve_user(email)
        if not user:
            from app.models.people import User

            user = User(
                name=str(name).strip(),
                email=str(email).strip(),
                team=str(ws.cell(row_idx, 11).value or "Shared"),
                role=str(ws.cell(row_idx, 7).value or "Operations"),
                allocation_percent=int(ws.cell(row_idx, 8).value or 100),
                weekly_hours=int(ws.cell(row_idx, 9).value or 40),
                active=True,
            )
            ctx.db.add(user)
            ctx.db.flush()
            ctx.cache_user(user)
        else:
            user.team = str(ws.cell(row_idx, 11).value or user.team)
            user.role = str(ws.cell(row_idx, 7).value or user.role)
            user.allocation_percent = int(ws.cell(row_idx, 8).value or user.allocation_percent)
            user.weekly_hours = int(ws.cell(row_idx, 9).value or user.weekly_hours)
            ctx.cache_user(user)


def _infer_workstream_for_column(header_label: str) -> str:
    label = normalize_header(header_label)
    if any(token in label for token in ["wireless", "ssid", "wlan", "rf design", "opt 43"]):
        return "WIRELESS"
    if any(token in label for token in ["sd-wan", "wan", "vpn", "dhcp", "dns", "arp", "circuit"]):
        return "SDWAN"
    if any(token in label for token in ["sda", "ip pool", "forescout", "end point"]):
        return "SDA"
    return "PROGRAM"


def _friendly_column_title(ws, header_row: int, col_idx: int) -> str:
    section_labels = []
    for row_idx in range(1, header_row):
        value = ws.cell(row_idx, col_idx).value
        if value not in (None, ""):
            section_labels.append(str(value).strip())
    leaf = str(ws.cell(header_row, col_idx).value or "").strip()
    parts = [part for part in section_labels + [leaf] if part]
    deduped = []
    for part in parts:
        if not deduped or deduped[-1] != part:
            deduped.append(part)
    return " / ".join(deduped) if deduped else f"Column {col_idx}"


def _column_owner_lookup(headers: dict[str, int]) -> dict[int, int]:
    owner_lookup: dict[int, int] = {}
    for header_name, header_col in headers.items():
        if header_name.endswith(" owner"):
            base_name = header_name.removesuffix(" owner")
            target_col = headers.get(base_name)
            if target_col:
                owner_lookup[target_col] = header_col
    return owner_lookup


def parse_upcoming_migrations_sheet(ctx: ImportContext, ws) -> None:
    header_row, normalized_headers = find_header_row(ws, {"site name", "schedule dat (et)"}, max_scan_rows=6)
    if header_row is None:
        header_row, normalized_headers = find_header_row(ws, {"site name", "schedule date (et)"}, max_scan_rows=6)
    if header_row is None:
        ctx.issue(ws.title, severity="error", message="Could not resolve site/date columns in upcoming migrations sheet.")
        return

    actual_headers = {
        col_idx: str(ws.cell(header_row, col_idx).value).strip()
        for col_idx in range(1, ws.max_column + 1)
        if ws.cell(header_row, col_idx).value not in (None, "")
    }
    site_col = normalized_headers.get("site name")
    date_col = normalized_headers.get("schedule date (et)") or normalized_headers.get("schedule dat (et)")
    site_type_col = normalized_headers.get("site type")
    if not site_col or not date_col:
        ctx.issue(ws.title, severity="error", message="Could not resolve site/date columns in upcoming migrations sheet.")
        return

    ignored_headers = {
        "site name",
        "schedule date (et)",
        "schedule dat (et)",
        "site type",
    }
    owner_lookup = _column_owner_lookup(normalized_headers)
    stakeholder_lookup = {
        target_col: stakeholder_col
        for header_name, stakeholder_col in normalized_headers.items()
        if header_name.endswith(" stakeholder") and (target_col := normalized_headers.get(header_name.removesuffix(" stakeholder")))
    }

    tracked_columns = [
        col_idx
        for header_name, col_idx in normalized_headers.items()
        if header_name not in ignored_headers and not header_name.endswith(" owner") and not header_name.endswith(" stakeholder")
    ]

    for row_idx in range(header_row + 1, ws.max_row + 1):
        site_name = ws.cell(row_idx, site_col).value
        if not site_name or str(site_name).startswith("Team"):
            continue
        scheduled_date = parse_date_cell(ws.cell(row_idx, date_col).value)
        site = ctx.resolve_site(site_name, market="Imported", wave=str(ws.cell(row_idx, site_type_col).value or "") or None)
        site.notes = f"Imported from {ws.title}"
        window = ctx.resolve_window(site, scheduled_date)
        if window:
            window.notes = f"Imported from {ws.title}"
        raw_row = {header_name: ws.cell(row_idx, col_idx).value for col_idx, header_name in actual_headers.items()}
        for col_idx in tracked_columns:
            raw_value = ws.cell(row_idx, col_idx).value
            owner_hint = ws.cell(row_idx, owner_lookup[col_idx]).value if col_idx in owner_lookup else None
            stakeholder_hint = ws.cell(row_idx, stakeholder_lookup[col_idx]).value if col_idx in stakeholder_lookup else None
            status = infer_status(raw_value, comment=raw_value, owner=owner_hint, aux=stakeholder_hint)
            if status is None:
                continue
            column_name = actual_headers[col_idx]
            title = _friendly_column_title(ws, header_row, col_idx)
            workstream_key = _infer_workstream_for_column(title)
            phase = "Readiness" if ws.title == "Upcoming Migrations" else "Implementation Plan"
            priority = PriorityLevel.HIGH if any(token in normalize_header(title) for token in ["checklist", "peer review", "wan circuits", "ip pools", "configuration scripts", "pre-staging", "rf design"]) else PriorityLevel.MEDIUM
            owner = ctx.resolve_user(owner_hint) or ctx.resolve_user(stakeholder_hint)
            ctx.upsert_task(
                site=site,
                workstream=ctx.get_workstream(workstream_key),
                title=title,
                phase=phase,
                description=str(raw_value).strip() if raw_value not in (None, "") else f"{column_name} imported from workbook planning sheet.",
                status=status,
                priority=priority,
                owner=owner,
                due_date=scheduled_date,
                source_tab=ws.title,
                source_row_key=f"{ws.title}:row-{row_idx}",
                source_column_key=title,
                migration_window_id=window.id if window else None,
                raw_import_json=raw_row,
            )


def _four_column_blocks(ws, header_row: int, start_col: int) -> list[dict]:
    blocks = []
    col = start_col
    while col <= ws.max_column:
        site_header = ws.cell(header_row, col).value
        if is_site_like(site_header):
            header_label = str(site_header).strip()
            blocks.append(
                {
                    "site_name": normalize_site_name(site_header),
                    "header_label": header_label,
                    "status_col": col,
                    "comment_col": col + 1 if col + 1 <= ws.max_column else None,
                    "owner_col": col + 2 if col + 2 <= ws.max_column else None,
                    "aux_col": col + 3 if col + 3 <= ws.max_column else None,
                }
            )
            col += 4
        else:
            col += 1
    return blocks


def _sdwan_blocks(ws) -> list[dict]:
    def site_from_value(raw_value: Any) -> tuple[str, str, str | None] | None:
        header = normalize_header(raw_value)
        if header.startswith("team:") or header.startswith("date:"):
            return None
        candidate = normalize_site_name(raw_value)
        if not is_site_like(candidate):
            return None
        label = str(raw_value).strip()
        engineer_match = re.search(r"\(([^)]+)\)", label)
        engineer_hint = engineer_match.group(1).strip() if engineer_match else None
        return candidate, label, engineer_hint

    def nearest_site_info(status_col: int) -> tuple[str, str, str | None] | None:
        for candidate_col in [status_col + 1, status_col - 1, status_col, status_col + 2, status_col - 2]:
            if 1 <= candidate_col <= ws.max_column:
                candidate = site_from_value(ws.cell(2, candidate_col).value)
                if candidate:
                    return candidate
        for candidate_col in [status_col, status_col + 1, status_col - 1, status_col + 2, status_col - 2]:
            if 1 <= candidate_col <= ws.max_column:
                candidate = site_from_value(ws.cell(1, candidate_col).value)
                if candidate:
                    return candidate
        return None

    blocks = []
    col = 3
    seen: set[tuple[str, int]] = set()
    while col <= ws.max_column:
        header2 = ws.cell(2, col).value
        header2_norm = normalize_header(header2)
        if is_site_like(header2) and normalize_header(ws.cell(2, col + 1).value) == "time":
            site_name = normalize_site_name(header2)
            header_label = str(header2).strip()
            engineer_match = re.search(r"\(([^)]+)\)", header_label)
            engineer_hint = engineer_match.group(1).strip() if engineer_match else None
            if (site_name, col) not in seen:
                seen.add((site_name, col))
                blocks.append(
                    {
                        "site_name": site_name,
                        "header_label": header_label,
                        "site_engineer_hint": engineer_hint,
                        "status_col": col,
                        "comment_col": None,
                        "owner_col": None,
                        "aux_col": col + 1,
                    }
                )
            col += 2
            continue
        if header2_norm == "status":
            site_info = nearest_site_info(col)
            if site_info and (site_info[0], col) not in seen:
                site_name, header_label, engineer_hint = site_info
                seen.add((site_name, col))
                aux_col = None
                for candidate_col in [col - 1, col + 1, col + 2]:
                    if 1 <= candidate_col <= ws.max_column and normalize_header(ws.cell(2, candidate_col).value) in {"time", "t-minus"}:
                        aux_col = candidate_col
                        break
                comment_col = col + 1 if col + 1 <= ws.max_column and normalize_header(ws.cell(2, col + 1).value) == "comments" else None
                blocks.append(
                    {
                        "site_name": site_name,
                        "header_label": header_label,
                        "site_engineer_hint": engineer_hint,
                        "status_col": col,
                        "comment_col": comment_col,
                        "owner_col": None,
                        "aux_col": aux_col,
                    }
                )
        col += 1
    return blocks


def _sdwan_scheduled_date(ws, status_col: int) -> date | None:
    for candidate_col in [status_col, status_col - 1, status_col + 1, status_col - 2, status_col + 2]:
        if 1 <= candidate_col <= ws.max_column:
            parsed = parse_date_cell(ws.cell(1, candidate_col).value)
            if parsed:
                return parsed
    return None


def _is_peer_review_label(value: Any) -> bool:
    return "peer review" in normalize_header(value)


def _split_tracker_status_detail(raw_status: Any) -> tuple[Any, str | None]:
    if raw_status in (None, ""):
        return raw_status, None
    text = str(raw_status).strip()
    match = re.match(
        r"^\s*(partial(?:ly)?\s+completed?|partial\s+complete|completed?|done|pending(?:\s*/\s*on\s*hold)?|not\s+started|in\s+progress|n\s*/?\s*a)\s*[-:]\s*(.+?)\s*$",
        text,
        re.IGNORECASE,
    )
    if not match:
        return text, None
    return match.group(1), match.group(2).strip()


def parse_tracker_sheet(ctx: ImportContext, ws) -> None:
    workstream = ctx.get_workstream(_workstream_key_for_sheet(ws.title))
    if ws.title == "SDA Tracker":
        header_row = 3
        data_start = 4
        blocks = _four_column_blocks(ws, header_row, 3)
    elif ws.title == "Wireless Tracker":
        header_row = 2
        data_start = 3
        blocks = _four_column_blocks(ws, header_row, 4)
    else:
        header_row = 2
        data_start = 3
        blocks = _sdwan_blocks(ws)

    current_phase = None
    for row_idx in range(data_start, ws.max_row + 1):
        phase_value = ws.cell(row_idx, 1).value
        task_title = ws.cell(row_idx, 2).value
        phase_text = str(phase_value).strip() if phase_value not in (None, "") else None
        title_text = str(task_title).strip() if task_title not in (None, "") else None

        wireless_reference_row = ws.title == "Wireless Tracker" and is_wireless_reference_row(phase_text, title_text)

        synthesized_peer_review_row = bool(phase_text and not title_text and _is_peer_review_label(phase_text))

        if phase_text and not synthesized_peer_review_row:
            current_phase = phase_text

        phase_for_row = "Peer Review" if synthesized_peer_review_row else (phase_text or current_phase)
        task_title = phase_text if synthesized_peer_review_row else title_text
        peer_review_row = synthesized_peer_review_row or _is_peer_review_label(phase_for_row) or _is_peer_review_label(task_title)

        if task_title in (None, ""):
            continue
        for block in blocks:
            status_value = ws.cell(row_idx, block["status_col"]).value
            status_seed, status_detail = _split_tracker_status_detail(status_value)
            comment_value = ws.cell(row_idx, block["comment_col"]).value if block["comment_col"] else None
            owner_value = ws.cell(row_idx, block["owner_col"]).value if block["owner_col"] else None
            aux_value = ws.cell(row_idx, block["aux_col"]).value if block["aux_col"] else None
            description_value = comment_value if comment_value not in (None, "") else status_detail
            status = infer_status(status_seed, comment=description_value, owner=owner_value, aux=aux_value) or TaskStatus.NOT_STARTED
            site = ctx.resolve_site(block["site_name"], market="Imported")
            scheduled_date = None
            due_date = None
            window = None
            scheduled_date = (
                None
                if wireless_reference_row
                else (
                    _sdwan_scheduled_date(ws, block["status_col"])
                    if ws.title == "SD-WAN Tracker"
                    else parse_date_cell(ws.cell(1, block["status_col"]).value)
                )
            )
            if scheduled_date and not wireless_reference_row:
                window = ctx.resolve_window(site, scheduled_date)
                due_date = parse_tminus_due(aux_value, scheduled_date)
            owner = None if wireless_reference_row else (ctx.resolve_user(owner_value) or ctx.resolve_user(block.get("site_engineer_hint")))
            ctx.upsert_task(
                site=site,
                workstream=workstream,
                title=task_title,
                phase=phase_for_row,
                description=str(description_value).strip() if description_value not in (None, "") else None,
                status=status,
                priority=PriorityLevel.HIGH if peer_review_row else PriorityLevel.MEDIUM,
                owner=owner,
                due_date=due_date,
                source_tab=ws.title,
                source_row_key=f"{ws.title}:row-{row_idx}",
                source_column_key=f"{block.get('header_label') or block['site_name']} ({get_column_letter(block['status_col'])})",
                migration_window_id=window.id if window else None,
                raw_import_json={
                    "sheet": ws.title,
                    "workstream_key": workstream.key,
                    "site_name": block["site_name"],
                    "header_label": block.get("header_label"),
                    "site_engineer_hint": block.get("site_engineer_hint"),
                    "phase": phase_for_row,
                    "task": task_title,
                    "status": status_seed,
                    "status_detail": status_detail,
                    "comment": description_value,
                    "owner": owner_value,
                    "aux": aux_value,
                    "peer_review_row": peer_review_row,
                    "reference_only_row": wireless_reference_row,
                    "scheduled_date": scheduled_date.isoformat() if scheduled_date else None,
                },
            )


def parse_pto_sheet(ctx: ImportContext, ws) -> None:
    for row_idx in range(2, ws.max_row + 1):
        name = ws.cell(row_idx, 1).value
        pto_range = ws.cell(row_idx, 2).value
        if not name or not pto_range:
            continue
        user = ctx.resolve_user(name)
        if not user:
            ctx.issue(ws.title, severity="warning", message=f"Could not match PTO owner '{name}' to an existing user.", row_number=row_idx, raw_row_json={"name": name, "pto": pto_range})
            continue
        start_date, end_date = parse_pto_range(pto_range)
        if not start_date or not end_date:
            ctx.issue(ws.title, severity="warning", message="Could not parse PTO range.", row_number=row_idx, field_name="PTO", raw_row_json={"name": name, "pto": pto_range})
            continue
        existing = ctx.db.scalar(
            select(PTO).where(PTO.user_id == user.id, PTO.start_date == start_date, PTO.end_date == end_date)
        )
        if existing:
            continue
        ctx.db.add(
            PTO(
                user_id=user.id,
                start_date=start_date,
                end_date=end_date,
                pto_type=PTOType.PTO,
                notes=str(ws.cell(row_idx, 3).value or ""),
                created_at=datetime.now(timezone.utc),
            )
        )


def parse_calls_sheet(ctx: ImportContext, ws) -> None:
    weekdays = []
    for col in range(2, ws.max_column + 1):
        header = ws.cell(1, col).value
        if header:
            weekdays.append((col, str(header).strip()))
    today = date.today()
    monday = today - timedelta(days=today.weekday())
    weekday_map = {"monday": 0, "tuesday": 1, "wednesday": 2, "thursday": 3, "friday": 4}
    for row_idx in range(2, ws.max_row + 1):
        slot = ws.cell(row_idx, 1).value
        if not slot:
            continue
        start_hour = 8
        try:
            start_hour = int(str(slot).split(":")[0].strip())
        except ValueError:
            ctx.issue(ws.title, severity="warning", row=row_idx, message=f"Could not parse time slot '{slot}', defaulting to 08:00")
        for col, weekday in weekdays:
            value = ws.cell(row_idx, col).value
            if not value:
                continue
            call_date = monday + timedelta(days=weekday_map.get(weekday.lower(), 0))
            ctx.db.add(
                Call(
                    title=str(value).strip(),
                    owner_id=None,
                    start_datetime=datetime.combine(call_date, time(start_hour % 24, 0), tzinfo=timezone.utc),
                    end_datetime=datetime.combine(call_date, time((start_hour + 1) % 24, 0), tzinfo=timezone.utc),
                    call_type=CallType.PLANNING,
                    related_site_id=None,
                    notes=f"Imported from {ws.title}",
                    created_at=datetime.now(timezone.utc),
                )
            )


def parse_poc_calendar_sheet(ctx: ImportContext, ws) -> None:
    for row_idx in range(7, ws.max_row + 1):
        team = ws.cell(row_idx, 10).value
        task_name = ws.cell(row_idx, 13).value
        comment = ws.cell(row_idx, 15).value
        if not team or not task_name:
            continue
        ctx.issue(
            ws.title,
            severity="warning",
            message="POC Calendar rows are preserved as traceability notes in phase 1; direct linkage is deferred.",
            row_number=row_idx,
            field_name="TASK",
            raw_row_json={"team": team, "task": task_name, "comment": comment},
        )


def import_workbook(db: Session, *, source: str | Path | bytes, file_name: str | None = None, imported_by: int | None = None) -> Import:
    from app.imports.parser_registry import PARSER_REGISTRY
    from app.services.sites import ensure_tracker_coverage

    record = Import(
        file_name=file_name or (Path(source).name if isinstance(source, (str, Path)) else "uploaded-workbook.xlsx"),
        imported_at=datetime.now(timezone.utc),
        imported_by=imported_by,
        source_type="excel_workbook",
        status=ImportStatus.PENDING,
        summary_json=None,
        created_tasks=0,
        updated_tasks=0,
        warnings_count=0,
        errors_count=0,
    )
    db.add(record)
    db.flush()

    workbook = load_workbook(filename=source if isinstance(source, (str, Path)) else BytesIO(source), data_only=True)
    ctx = ImportContext(db=db, import_record=record)
    processed = []
    for sheet_name in workbook.sheetnames:
        parser = PARSER_REGISTRY.get(sheet_name)
        if not parser:
            continue
        processed.append(sheet_name)
        parser(ctx, workbook[sheet_name])

    tracker_tabs = {"SD-WAN Tracker", "SDA Tracker", "Wireless Tracker"}
    pruned_tracker_tasks = ctx.prune_stale_tracker_tasks(source_tabs=set(processed) & tracker_tabs)

    coverage = ensure_tracker_coverage(
        db,
        commit=False,
        source_type="tracker_backfill",
        created_from="import_coverage_repair",
    )

    record.created_tasks = ctx.created_tasks + coverage["tasks_created"]
    record.updated_tasks = ctx.updated_tasks
    record.warnings_count = ctx.warnings
    record.errors_count = ctx.errors
    record.summary_json = {
        "processed_sheets": processed,
        "available_sheets": workbook.sheetnames,
        "pruned_tracker_tasks": pruned_tracker_tasks,
        "coverage_backfill": coverage,
    }
    record.status = ImportStatus.SUCCESS
    if ctx.errors:
        record.status = ImportStatus.PARTIAL if processed else ImportStatus.FAILED
    db.commit()
    db.refresh(record)
    return record
