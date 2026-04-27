from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from difflib import SequenceMatcher
from typing import Any
import unicodedata

from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.core.timeline import timeline_due_date_from_text
from app.core.enums import ImportStatus, PriorityLevel, TaskStatus, normalize_status
from app.models.execution import Artifact, Task, TaskUpdate
from app.models.imports import Import, ImportError
from app.models.people import Assignment, User
from app.models.planning import MigrationWindow, Site, Workstream


GENERIC_HEADERS = {"phase", "task", "comments", "comment", "engineer", "time", "t-minus", "status", "helpful links"}

WORKSTREAM_DEFAULTS: dict[str, tuple[str, str]] = {
    "PROGRAM": ("Program", "Cross-workstream readiness, plans, and governance"),
    "SDWAN": ("SD-WAN", "SD-WAN circuit, template, and cutover preparation"),
    "SDA": ("SDA", "SDA implementation and validation"),
    "WIRELESS": ("Wireless", "RF design, enclosure validation, and wireless cutover"),
    "CUTOVER": ("Cutover", "Execution night, go/no-go, and bridge readiness"),
    "DAY1": ("Day 1", "Post-migration support and issue handling"),
}

WORKSTREAM_ALIASES: dict[str, str] = {
    "SD-WAN": "SDWAN",
    "SD WAN": "SDWAN",
    "DAY 1": "DAY1",
}


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text.strip().lower())


def canonical_workstream_key(value: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        return "PROGRAM"
    direct = WORKSTREAM_ALIASES.get(raw.upper())
    if direct:
        return direct
    collapsed = re.sub(r"[^A-Z0-9]+", "", raw.upper())
    if collapsed in WORKSTREAM_DEFAULTS:
        return collapsed
    if collapsed in {"SDW", "SDWANTRACKER"}:
        return "SDWAN"
    if collapsed in {"WIRELESSTRACKER", "WIFI"}:
        return "WIRELESS"
    if collapsed in {"SDATRACKER"}:
        return "SDA"
    return collapsed or "PROGRAM"


def normalize_site_name(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^\s*site:\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"^\s*team:\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\(.*?\)", "", text)
    text = re.sub(r"\s*-\s*\d{1,2}/\d{1,2}/\d{2,4}.*$", "", text)
    text = re.sub(r"^\s*date:\s*.*$", "", text, flags=re.IGNORECASE)
    text = text.split(",")[0]
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ,")


def canonical_site_identity(value: Any) -> str:
    text = normalize_site_name(value).lower()
    text = re.sub(r"\b(street|st|road|rd|drive|dr|avenue|ave|way|parkway|pkwy|square|sq)\b\.?$", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ,")


def site_slug(value: Any) -> str:
    normalized = canonical_site_identity(value)
    normalized = normalized.replace("&", "and")
    normalized = re.sub(r"[^a-z0-9]+", "-", normalized)
    return normalized.strip("-")


def is_site_like(value: Any) -> bool:
    text = normalize_header(value)
    if not text or text in GENERIC_HEADERS:
        return False
    if len(text) <= 2:
        return False
    return any(char.isdigit() for char in text) or any(token in text for token in ["street", "ave", "road", "park", "way", "square", "main", "mount", "central", "washington"])


def parse_tminus_due(value: Any, scheduled_date: date | None) -> date | None:
    if scheduled_date is None or value in (None, ""):
        return None
    return timeline_due_date_from_text(str(value).strip(), scheduled_date)


def infer_status(raw_status: Any, *, comment: Any = None, owner: Any = None, aux: Any = None) -> TaskStatus | None:
    if raw_status not in (None, ""):
        text = str(raw_status).strip()
        normalized = normalize_status(text)
        if normalized != TaskStatus.NOT_STARTED or text.lower().strip() in {"pending", "not started", "", "blank"}:
            return normalized
        if text:
            return TaskStatus.IN_PROGRESS
    if any(item not in (None, "") for item in [comment, owner, aux]):
        return TaskStatus.IN_PROGRESS
    return None


def parse_date_cell(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        value = re.sub(r"^\s*date:\s*", "", value, flags=re.IGNORECASE)
        value = re.sub(r"(\d{1,2})(st|nd|rd|th)\b", r"\1", value, flags=re.IGNORECASE)
        value = re.sub(r",", "", value)
        value = re.sub(r"\s+", " ", value).strip()
        embedded = re.search(r"\b\d{1,2}/\d{1,2}/\d{2,4}\b", value)
        if embedded:
            value = embedded.group(0)
        for fmt in (
            "%Y-%m-%d",
            "%m/%d/%Y",
            "%m/%d/%y",
            "%d-%b-%y",
            "%d-%b-%Y",
            "%d %b %y",
            "%d %b %Y",
            "%b %d %Y",
            "%B %d %Y",
            "%b %d",
            "%B %d",
        ):
            try:
                parsed = datetime.strptime(value, fmt).date()
                if "%Y" not in fmt and "%y" not in fmt:
                    parsed = parsed.replace(year=date.today().year)
                return parsed
            except ValueError:
                continue
    return None


MONTHS = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}


def parse_pto_range(value: Any, anchor_year: int | None = None) -> tuple[date | None, date | None]:
    if value in (None, ""):
        return None, None
    text = str(value).strip()
    anchor_year = anchor_year or date.today().year

    compact = re.sub(r"\s+", "", text)
    patterns = [
        re.match(r"([A-Za-z]{3})(\d{1,2})-([A-Za-z]{3})?(\d{1,2})", compact),
        re.match(r"(\d{1,2})([A-Za-z]{3})-?(\d{1,2})([A-Za-z]{3})", compact),
    ]
    match = next((item for item in patterns if item), None)

    if match and match.re.pattern.startswith("([A-Za-z]"):
        start_month = MONTHS.get(match.group(1)[:3].lower())
        start_day = int(match.group(2))
        end_month = MONTHS.get((match.group(3) or match.group(1))[:3].lower())
        end_day = int(match.group(4))
    elif match:
        start_day = int(match.group(1))
        start_month = MONTHS.get(match.group(2)[:3].lower())
        end_day = int(match.group(3))
        end_month = MONTHS.get(match.group(4)[:3].lower())
    else:
        alt = re.match(r"(\d{1,2})to([A-Za-z]{3})(\d{1,2})", compact, re.IGNORECASE)
        if not alt:
            return None, None
        start_day = int(alt.group(1))
        end_month = MONTHS.get(alt.group(2)[:3].lower())
        end_day = int(alt.group(3))
        if not end_month:
            return None, None
        start_month = 12 if end_month == 1 else end_month - 1

    if not start_month or not end_month:
        return None, None
    start_year = anchor_year
    end_year = anchor_year
    if end_month < start_month:
        end_year += 1
    return date(start_year, start_month, start_day), date(end_year, end_month, end_day)


def find_header_row(ws, required_headers: set[str], max_scan_rows: int = 8) -> tuple[int | None, dict[str, int]]:
    for row_idx in range(1, min(ws.max_row, max_scan_rows) + 1):
        headers = {
            normalize_header(ws.cell(row_idx, col_idx).value): col_idx
            for col_idx in range(1, ws.max_column + 1)
            if ws.cell(row_idx, col_idx).value not in (None, "")
        }
        if required_headers.issubset(set(headers.keys())):
            return row_idx, headers
    return None, {}


def to_jsonable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(key): to_jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [to_jsonable(item) for item in value]
    return value


@dataclass
class ImportContext:
    db: Session
    import_record: Import
    site_cache: dict[str, Site] = field(default_factory=dict)
    user_cache: dict[str, User] = field(default_factory=dict)
    workstream_cache: dict[str, Workstream] = field(default_factory=dict)
    task_cache: dict[tuple[int, int | None, str | None, str], Task] = field(default_factory=dict)
    seen_import_task_ids: set[int] = field(default_factory=set)
    warnings: int = 0
    errors: int = 0
    created_tasks: int = 0
    updated_tasks: int = 0

    def __post_init__(self):
        for site in self.db.scalars(select(Site)).all():
            self.site_cache[site_slug(site.site_name)] = site
        for user in self.db.scalars(select(User)).all():
            self.cache_user(user)
        for workstream in self.db.scalars(select(Workstream)).all():
            canonical_key = canonical_workstream_key(workstream.key)
            self.workstream_cache[canonical_key] = workstream
            self.workstream_cache[workstream.key] = workstream
        for task in self.db.scalars(select(Task)).all():
            self.task_cache[(task.site_id, task.workstream_id, task.phase, task.title)] = task

    def cache_user(self, user: User) -> None:
        name_key = normalize_header(user.name)
        self.user_cache[name_key] = user
        if user.name:
            self.user_cache[normalize_header(user.name.split()[0])] = user
        if user.email:
            self.user_cache[normalize_header(user.email)] = user

    def issue(self, sheet_name: str, *, severity: str, message: str, row_number: int | None = None, field_name: str | None = None, raw_row_json: dict | None = None) -> None:
        self.db.add(
            ImportError(
                import_id=self.import_record.id,
                sheet_name=sheet_name,
                row_number=row_number,
                field_name=field_name,
                severity=severity,
                message=message,
                raw_row_json=to_jsonable(raw_row_json),
            )
        )
        if severity.lower() == "warning":
            self.warnings += 1
        else:
            self.errors += 1

    def get_workstream(self, key: str) -> Workstream:
        canonical_key = canonical_workstream_key(key)
        cached = self.workstream_cache.get(canonical_key) or self.workstream_cache.get(key)
        if cached:
            return cached

        existing = self.db.scalar(select(Workstream).where(Workstream.key == canonical_key))
        if existing:
            self.workstream_cache[canonical_key] = existing
            self.workstream_cache[existing.key] = existing
            return existing

        display_name, description = WORKSTREAM_DEFAULTS.get(
            canonical_key,
            (canonical_key.replace("_", " ").title(), "Created during workbook import"),
        )
        created = Workstream(key=canonical_key, name=display_name, description=description)
        self.db.add(created)
        self.db.flush()
        self.workstream_cache[canonical_key] = created
        self.workstream_cache[created.key] = created
        return created

    def resolve_user(self, raw_value: Any) -> User | None:
        if raw_value in (None, ""):
            return None
        key = normalize_header(raw_value)
        user = self.user_cache.get(key) or self.user_cache.get(key.split(" ")[0])
        if user:
            return user
        stmt = select(User).where((User.email == str(raw_value).strip()) | (User.name == str(raw_value).strip()))
        user = self.db.scalar(stmt)
        if user:
            self.cache_user(user)
        return user

    def resolve_site(self, raw_value: Any, *, region: str = "Northeast", market: str = "Unassigned", wave: str | None = None) -> Site:
        slug = site_slug(raw_value)
        if slug in self.site_cache:
            return self.site_cache[slug]
        identity = canonical_site_identity(raw_value)
        identity_match = re.match(r"^(\d+)\s+(.*)$", identity)
        best_site = None
        best_score = 0.0
        for existing_site in self.site_cache.values():
            existing_identity = canonical_site_identity(existing_site.site_name)
            existing_match = re.match(r"^(\d+)\s+(.*)$", existing_identity)
            if identity_match and existing_match and identity_match.group(1) != existing_match.group(1):
                continue
            score = SequenceMatcher(None, identity, existing_identity).ratio()
            if score > best_score:
                best_score = score
                best_site = existing_site
        if best_site and best_score >= 0.9:
            self.site_cache[slug] = best_site
            return best_site
        site = Site(
            site_code=f"IMP-{slug[:24].upper()}",
            site_name=normalize_site_name(raw_value),
            region=region,
            market=market,
            migration_wave=wave,
            address=None,
            notes="Created during workbook import",
            active=True,
        )
        self.db.add(site)
        self.db.flush()
        self.site_cache[slug] = site
        return site

    def resolve_window(self, site: Site, scheduled_date: date | None) -> MigrationWindow | None:
        if scheduled_date is None:
            return None
        window = self.db.scalar(
            select(MigrationWindow).where(
                MigrationWindow.site_id == site.id,
                MigrationWindow.scheduled_date == scheduled_date,
            )
        )
        if window:
            return window
        window = MigrationWindow(site_id=site.id, scheduled_date=scheduled_date, status="PLANNED")
        self.db.add(window)
        self.db.flush()
        return window

    def upsert_task(
        self,
        *,
        site: Site,
        workstream: Workstream,
        title: str,
        phase: str | None,
        description: str | None,
        status: TaskStatus,
        priority: PriorityLevel = PriorityLevel.MEDIUM,
        owner: User | None = None,
        due_date: date | None = None,
        source_type: str = "excel_workbook",
        source_tab: str | None = None,
        source_row_key: str | None = None,
        source_column_key: str | None = None,
        migration_window_id: int | None = None,
        raw_import_json: dict | None = None,
    ) -> Task:
        cache_key = (site.id, workstream.id, phase, title)
        task = self.task_cache.get(cache_key)
        if task is None:
            task = Task(
                site_id=site.id,
                workstream_id=workstream.id,
                migration_window_id=migration_window_id,
                phase=phase,
                title=title,
                description=description,
                status=status,
                priority=priority,
                owner_id=owner.id if owner else None,
                due_date=due_date,
                source_type=source_type,
                source_tab=source_tab,
                source_row_key=source_row_key,
                source_column_key=source_column_key,
                raw_import_json=to_jsonable(raw_import_json),
            )
            self.db.add(task)
            self.db.flush()
            self.task_cache[cache_key] = task
            self.seen_import_task_ids.add(task.id)
            self.created_tasks += 1
            return task

        task.description = description or task.description
        task.status = status
        task.priority = priority
        task.owner_id = owner.id if owner else task.owner_id
        task.due_date = due_date or task.due_date
        task.migration_window_id = migration_window_id or task.migration_window_id
        task.source_type = source_type
        task.source_tab = source_tab or task.source_tab
        task.source_row_key = source_row_key or task.source_row_key
        task.source_column_key = source_column_key or task.source_column_key
        task.raw_import_json = to_jsonable(raw_import_json) or task.raw_import_json
        self.seen_import_task_ids.add(task.id)
        self.updated_tasks += 1
        return task

    def prune_stale_tracker_tasks(self, *, source_tabs: set[str]) -> int:
        if not source_tabs:
            return 0

        stale_tasks = self.db.scalars(
            select(Task)
            .options(
                selectinload(Task.updates),
                selectinload(Task.assignments),
                selectinload(Task.artifacts),
            )
            .where(
                Task.source_type == "excel_workbook",
                Task.source_tab.in_(source_tabs),
            )
        ).unique().all()

        deleted = 0
        for task in stale_tasks:
            if task.id in self.seen_import_task_ids:
                continue

            for update in list(task.updates):
                self.db.delete(update)
            for assignment in list(task.assignments):
                self.db.delete(assignment)
            for artifact in list(task.artifacts):
                self.db.delete(artifact)

            self.task_cache.pop((task.site_id, task.workstream_id, task.phase, task.title), None)
            self.db.delete(task)
            deleted += 1

        return deleted
