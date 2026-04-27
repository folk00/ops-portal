from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


HEADER_ROW_HINTS = {
    "Ops team": 2,
    "POC Calendar": 6,
    "Automation Use Cases": 2,
    "Pre-Staging Activities": 4,
    "Upcoming Migrations": 4,
    "SDA Tracker": 3,
    "MW Activities": 3,
    "Wireless Tracker": 2,
    "SD-WAN Tracker": 2,
    "Codes": 2,
    "Day 1 Tracker": 1,
    "Best Practices": 1,
    "Pavel's Tab": 3,
    "Ops Calls": 1,
    "PTOs": 1,
    "Medium T-Minus": 1,
}

SHEET_TYPES = {
    "Ops team": "roster",
    "POC Calendar": "coordination",
    "Automation Use Cases": "reference",
    "Pre-Staging Activities": "reference",
    "Upcoming Migrations": "migration-planning",
    "SDA Tracker": "tracker-matrix",
    "MW Activities": "reference",
    "Wireless Tracker": "tracker-matrix",
    "SD-WAN Tracker": "tracker-matrix",
    "Codes": "reference",
    "Day 1 Tracker": "day1-flat",
    "Best Practices": "reference",
    "Pavel's Tab": "migration-planning",
    "Ops Calls": "calendar-matrix",
    "PTOs": "pto-roster",
    "Medium T-Minus": "reference",
}


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def normalized_token(value: Any) -> str:
    return normalize_header(value).lower()


def find_header_row(ws) -> int:
    if ws.title in HEADER_ROW_HINTS:
        return HEADER_ROW_HINTS[ws.title]
    best_row = 1
    best_score = -1
    for row_idx in range(1, min(ws.max_row, 10) + 1):
        values = [normalize_header(ws.cell(row_idx, c).value) for c in range(1, ws.max_column + 1)]
        non_empty = [value for value in values if value]
        if not non_empty:
            continue
        score = sum(1 for value in non_empty if any(char.isalpha() for char in value)) - sum(value.isdigit() for value in non_empty)
        if score > best_score:
            best_row = row_idx
            best_score = score
    return best_row


def classify_sheet(sheet_name: str) -> str:
    return SHEET_TYPES.get(sheet_name, "other")


def find_likely_columns(headers: list[str]) -> dict[str, list[str]]:
    categories = {
        "site": ["site", "location", "market"],
        "owner": ["owner", "engineer", "assigned", "assignee", "team", "poc", "lead"],
        "status": ["status", "completed", "ongoing", "not applicable", "na"],
        "date": ["date", "start", "end", "schedule", "calendar", "week", "day", "t-minus"],
        "comments": ["comment", "note", "remark", "challenge", "escalation", "issue"],
        "workstream": ["sd-wan", "sda", "wireless", "voice", "technology", "phase"],
        "blockers": ["blocked", "hold", "issue", "challenge", "escalation", "dependency"],
    }
    findings: dict[str, list[str]] = {key: [] for key in categories}
    for header in headers:
        token = normalized_token(header)
        for group, keywords in categories.items():
            if any(keyword in token for keyword in keywords):
                findings[group].append(header)
    return findings


def repeated_fields(headers: list[str]) -> dict[str, int]:
    counts = Counter(headers)
    return {header: count for header, count in counts.items() if count > 1}


def summarize_sheet(ws) -> dict[str, Any]:
    header_row = find_header_row(ws)
    headers = [normalize_header(ws.cell(header_row, c).value) for c in range(1, ws.max_column + 1)]
    non_empty_headers = [header for header in headers if header]
    return {
        "name": ws.title,
        "kind": classify_sheet(ws.title),
        "header_row": header_row,
        "max_rows": ws.max_row,
        "max_cols": ws.max_column,
        "headers": non_empty_headers,
        "likely_columns": find_likely_columns(non_empty_headers),
        "repeated_fields": repeated_fields(non_empty_headers),
    }


def render_markdown(path: Path, summaries: list[dict[str, Any]]) -> str:
    lines = [
        "# Workbook Analysis",
        "",
        f"Source workbook: `{path}`",
        "",
        "This file is generated from `scripts/workbook_inspect.py` using openpyxl. It is used as a primary domain reference for schema and import design in `ops-ops-portal`.",
        "",
        "## Sheet Inventory",
        "",
        "| Sheet | Kind | Header row | Rows | Cols |",
        "| --- | --- | ---: | ---: | ---: |",
    ]
    for item in summaries:
        lines.append(f"| {item['name']} | {item['kind']} | {item['header_row']} | {item['max_rows']} | {item['max_cols']} |")

    lines.extend(
        [
            "",
            "## Key Findings",
            "",
            "- The workbook is a hybrid of matrix trackers, planning tables, calendars, roster data, and reference sheets.",
            "- `SDA Tracker`, `Wireless Tracker`, and `SD-WAN Tracker` are matrix-style sheets where rows are activities and repeated site blocks live across columns.",
            "- `Upcoming Migrations` and `Pavel's Tab` act as migration-planning sheets with schedule dates and readiness columns.",
            "- `Ops team` is the capacity roster source and should influence user/team/allocation modeling.",
            "- `PTOs` and `Ops Calls` are schedule overlays rather than task trackers.",
            "- `POC Calendar` behaves more like coordination metadata and issue traceability than a core task table.",
            "",
            "## Per-Sheet Detail",
            "",
        ]
    )

    for item in summaries:
        lines.extend(
            [
                f"### {item['name']}",
                "",
                f"- Kind: `{item['kind']}`",
                f"- Header row: `{item['header_row']}`",
                f"- Header candidates: {', '.join(f'`{header}`' for header in item['headers'][:24]) or 'None'}",
            ]
        )
        for key, values in item["likely_columns"].items():
            if values:
                lines.append(f"- Likely {key} fields: {', '.join(f'`{value}`' for value in values[:16])}")
        if item["repeated_fields"]:
            repeated = ", ".join(f"`{header}` x{count}" for header, count in list(item["repeated_fields"].items())[:12])
            lines.append(f"- Repeated headers: {repeated}")
        lines.append("")

    lines.extend(
        [
            "## Schema Implications",
            "",
            "- Tasks need both `source_row_key` and `source_column_key` because tracker provenance is row-plus-column, not row-only.",
            "- Users need capacity metadata (`allocation_percent`, `weekly_hours`) because the roster sheet contains real staffing context.",
            "- Migration windows must be modeled separately from tasks because schedule dates appear in planning sheets and matrix headers.",
            "- Raw imported rows must be stored in JSON/JSONB because the workbook contains rich comments and mixed-format task notes.",
            "- Status normalization must tolerate `Complete`, `Completed`, `Complete.`, `In progress`, `Pending / On hold`, `N/A`, and blanks.",
            "",
            "## Parser Priorities",
            "",
            "1. `Ops team` -> users and capacity metadata",
            "2. `Upcoming Migrations` / `Pavel's Tab` -> sites and migration windows",
            "3. `SDA Tracker` / `Wireless Tracker` / `SD-WAN Tracker` -> task import using matrix column blocks",
            "4. `PTOs` -> PTO overlays",
            "5. `Ops Calls` -> calendar events",
            "6. `POC Calendar` -> coordination notes and future traceability hooks",
            "",
        ]
    )
    return "\n".join(lines)


def resolve_workbook_path(path_arg: str | None) -> Path:
    candidates = []
    if path_arg:
        candidates.append(Path(path_arg))
    candidates.extend(
        [
            Path.cwd() / "workbook.xlsx",
            Path.cwd().parent / "workbook.xlsx",
            Path.home() / "Downloads" / "workbook.xlsx",
        ]
    )
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("Could not locate 'workbook.xlsx'. Pass the path explicitly.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Inspect workbook structure for Ops Portal import design.")
    parser.add_argument("workbook", nargs="?", help="Path to the workbook")
    parser.add_argument("--out", default=str(Path(__file__).resolve().parents[1] / "docs" / "workbook-analysis.md"))
    args = parser.parse_args()

    workbook_path = resolve_workbook_path(args.workbook)
    workbook = load_workbook(workbook_path, data_only=True)
    summaries = [summarize_sheet(ws) for ws in workbook.worksheets]
    markdown = render_markdown(workbook_path, summaries)
    output_path = Path(args.out)
    output_path.write_text(markdown, encoding="utf-8")
    print(f"Wrote workbook analysis to {output_path}")
    print("Sheets:")
    for item in summaries:
        print(f"- {item['name']} (header_row={item['header_row']}, kind={item['kind']})")


if __name__ == "__main__":
    main()
